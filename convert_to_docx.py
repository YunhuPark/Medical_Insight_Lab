"""
paper_draft.md → paper_draft.docx 변환
python convert_to_docx.py
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re

def md_to_docx(md_path, docx_path):
    doc = Document()

    # 페이지 여백 설정 (A4 기준)
    for section in doc.sections:
        section.page_width  = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin   = Cm(3)
        section.right_margin  = Cm(3)
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)

    # 기본 스타일
    style = doc.styles['Normal']
    style.font.name = '맑은 고딕'
    style.font.size = Pt(10)

    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_rows = []
    i = 0

    while i < len(lines):
        line = lines[i].rstrip('\n')

        # 빈 줄
        if not line.strip():
            if in_table:
                _flush_table(doc, table_rows)
                in_table = False
                table_rows = []
            i += 1
            continue

        # 수평선
        if line.strip().startswith('---'):
            doc.add_paragraph('─' * 60)
            i += 1
            continue

        # 표 행
        if line.strip().startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        else:
            if in_table:
                _flush_table(doc, table_rows)
                in_table = False
                table_rows = []

        # 제목 처리
        h_match = re.match(r'^(#{1,4})\s+(.*)', line)
        if h_match:
            level = len(h_match.group(1))
            text  = h_match.group(2)
            text  = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
            p = doc.add_heading(text, level=min(level, 4))
            p.runs[0].font.name = '맑은 고딕' if p.runs else None
            i += 1
            continue

        # 코드 블록
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            run = p.add_run('\n'.join(code_lines))
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            i += 1
            continue

        # 글머리 기호
        if line.strip().startswith('- '):
            text = line.strip()[2:]
            text = _clean_inline(text)
            p = doc.add_paragraph(style='List Bullet')
            _add_inline(p, text)
            i += 1
            continue

        # 일반 단락
        text = _clean_inline(line)
        p = doc.add_paragraph()
        _add_inline(p, text)
        i += 1

    if in_table:
        _flush_table(doc, table_rows)

    doc.save(docx_path)
    print(f"저장 완료: {docx_path}")


def _clean_inline(text):
    return text


def _add_inline(p, text):
    """**bold**, `code`, 일반 텍스트 인라인 처리"""
    pattern = re.compile(r'\*\*(.*?)\*\*|`(.*?)`|(.*?)(?=\*\*|`|$)', re.DOTALL)
    for m in pattern.finditer(text):
        bold, code, plain = m.group(1), m.group(2), m.group(3)
        if bold:
            run = p.add_run(bold)
            run.bold = True
            run.font.name = '맑은 고딕'
            run.font.size = Pt(10)
        elif code:
            run = p.add_run(code)
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        elif plain:
            run = p.add_run(plain)
            run.font.name = '맑은 고딕'
            run.font.size = Pt(10)


def _flush_table(doc, rows):
    # 구분선(---|---) 행 제거
    data_rows = [r for r in rows if not re.match(r'^\s*\|[-| :]+\|\s*$', r)]
    if not data_rows:
        return

    parsed = []
    for row in data_rows:
        cells = [c.strip() for c in row.strip().strip('|').split('|')]
        parsed.append(cells)

    if not parsed:
        return

    ncols = max(len(r) for r in parsed)
    table = doc.add_table(rows=len(parsed), cols=ncols)
    table.style = 'Table Grid'

    for r_idx, row in enumerate(parsed):
        for c_idx, cell_text in enumerate(row):
            if c_idx >= ncols:
                break
            cell = table.cell(r_idx, c_idx)
            cell.text = ''
            p = cell.paragraphs[0]
            # 헤더 행 굵게
            run = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text))
            run.font.name = '맑은 고딕'
            run.font.size = Pt(9)
            if r_idx == 0:
                run.bold = True


if __name__ == '__main__':
    md_to_docx('paper_draft.md', 'paper_draft.docx')
    md_to_docx('cover_letter.md', 'cover_letter.docx')
    print("논문 + 커버레터 DOCX 변환 완료")
